"""Requirement specs: the Markdown in specs/ is the source, the .xlsx workbooks are generated for viewing.

    python tools/specs.py build [--force]   write specs/xlsx/*.xlsx and the generated summary in specs/requirements/README.md
    python tools/specs.py diff              list what was changed by hand in the workbooks since the last build

One workbook per Markdown file: overview.xlsx (README.md: summary numbers as values, each mechanic linking to its
workbook), <mechanic>.xlsx, glossary.xlsx.

build validates the Markdown (enum values, duplicate or unknown IDs, table rows matching their sections) and refuses to
overwrite a workbook that was edited since the last build (hashes in specs/xlsx/.built.json): run diff, copy the changes
into the Markdown, then build --force. Output is byte-for-byte deterministic, so an unchanged spec leaves the files alone.

Markdown layout of a mechanic file (specs/requirements/<slug>.md):
    # <Sheet name>
    <pillar line>                                  paragraph
    <where it stands>                              paragraph
    ## Requirements
    | ID | Feature | Priority | ... |               short fields, one row per requirement
                                                   Built (Yes/Partial/No) = what the code does as of the date in
                                                   'where it stands'; must agree with Status (Done = Yes, In Progress
                                                   = Partial, anything else = No)
    ### <ID> <Feature>                             long fields: **Requirement:** **Acceptance criteria:** **As built:** **Notes:**
    ## Decisions & open questions
    | # | Question | Related IDs | ... |            short fields; # is OQ-<n>
    ### OQ-<n> <Question>                          long fields: **Options / notes:** **Decision:**
"""
import datetime
import hashlib
import io
import json
import re
import sys
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
SPECS = ROOT / "specs"
REQ_DIR = SPECS / "requirements"
XLSX_DIR = SPECS / "xlsx"
HASHES = XLSX_DIR / ".built.json"
OVERVIEW_BOOK = "overview.xlsx"
GLOSSARY_BOOK = "glossary.xlsx"

PRIORITY = ["Must", "Should", "Could", "Won't"]
STATUS = ["Proposed", "Approved", "In Progress", "Done", "Deferred", "Rejected"]
MILESTONE = ["MVP", "Alpha", "Beta", "Release", "Post-Launch"]
TYPE = ["Functional", "Non-Functional"]
FLAG = ["Conflict", "Review", "Updated", "New"]
STATE = ["Open", "Provisional", "Decided"]
BUILT = ["Yes", "Partial", "No"]
BUILT_FOR_STATUS = {"Done": "Yes", "In Progress": "Partial"}  # every other status: "No"
STATUS_ORDER = ["Done", "In Progress", "Approved", "Proposed", "Deferred", "Rejected"]  # Overview and summary order

# (key, workbook header, Markdown table header or None for a section field, Markdown section label, allowed values)
REQ_FIELDS = [
    ("id", "ID", "ID", None, None),
    ("feature", "Feature", "Feature", None, None),
    ("requirement", "Requirement", None, "Requirement", None),
    ("acceptance", "Acceptance Criteria", None, "Acceptance criteria", None),
    ("as_built", "As Built", None, "As built", None),
    ("built", "Built", "Built", None, BUILT),
    ("flag", "Audit Flag", "Audit Flag", None, FLAG),
    ("notes", "Notes", None, "Notes", None),
    ("priority", "Priority", "Priority", None, PRIORITY),
    ("status", "Status", "Status", None, STATUS),
    ("milestone", "Milestone", "Milestone", None, MILESTONE),
    ("type", "Type", "Type", None, TYPE),
    ("pillar", "Pillar", "Pillar", None, None),
    ("depends", "Depends On", "Depends On", None, None),
    ("owner", "Owner", "Owner", None, None),
]
REQ_TABLE = ["id", "feature", "priority", "status", "built", "milestone", "type", "pillar", "depends", "flag", "owner"]
Q_FIELDS = [
    ("num", "#", "#", None, None),
    ("related", "Related IDs", "Related IDs", None, None),
    ("question", "Question", "Question", None, None),
    ("options", "Options / Notes", None, "Options / notes", None),
    ("decision", "Decision", None, "Decision", None),
    ("state", "State", "State", None, STATE),
    ("decided_by", "Decided By", "Decided By", None, None),
    ("date", "Date", "Date", None, None),
]
Q_TABLE = ["num", "question", "related", "state", "decided_by", "date"]
ID_RE = re.compile(r"^[A-Z]{2}-\d{3}$")


@dataclass
class Sheet:
    name: str
    file: str = ""
    intro: list = field(default_factory=list)
    reqs: list = field(default_factory=list)
    questions: list = field(default_factory=list)


@dataclass
class Spec:
    title: str = ""
    intro: list = field(default_factory=list)
    sections: list = field(default_factory=list)  # (heading, [bullets]) shown on the Overview
    pillars: list = field(default_factory=list)  # (id, name, statement)
    sheets: list = field(default_factory=list)
    glossary_intro: list = field(default_factory=list)
    glossary: list = field(default_factory=list)  # (term, definition)


class SpecError(Exception):
    pass


def slug(text):
    """GitHub's heading anchor: lowercase, punctuation dropped, spaces to hyphens."""
    return re.sub(r"[^\w\- ]", "", text.lower()).replace(" ", "-")


def file_slug(name):
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")


# ---------------------------------------------------------------- Markdown -> model

def _blocks(text):
    """Splits a Markdown file into (level, heading, body lines) blocks; level 0 is the text before the first heading."""
    blocks, level, head, body = [], 0, "", []
    in_generated = False
    for line in text.splitlines():
        if line.startswith("<!-- BEGIN generated"):
            in_generated = True
        if in_generated:
            in_generated = not line.startswith("<!-- END generated")
            continue
        m = re.match(r"^(#{1,3}) (.*)$", line)
        if m:
            blocks.append((level, head, body))
            level, head, body = len(m.group(1)), m.group(2).strip(), []
        else:
            body.append(line)
    blocks.append((level, head, body))
    return blocks


def _paragraphs(lines):
    paras, cur = [], []
    for line in lines + [""]:
        if line.strip():
            cur.append(line.strip())
        elif cur:
            paras.append(" ".join(cur))
            cur = []
    return paras


def _split_row(line):
    cells = re.split(r"(?<!\\)\|", line.strip())[1:-1]
    return [re.sub(r"^\[(.*?)\]\(.*?\)$", r"\1", c.strip()).replace("\\|", "|") for c in cells]


def _table(lines):
    rows = [_split_row(l) for l in lines if l.strip().startswith("|")]
    if not rows:
        return None, []
    return rows[0], [r for r in rows[2:]]


def _fields(lines, labels, where):
    """Reads '**Label:** text' fields; a field runs until the next label, paragraph breaks kept as blank lines."""
    out, key, buf = {}, None, []

    def flush():
        if key:
            out[key] = "\n\n".join(_paragraphs(buf))

    for line in lines:
        m = re.match(r"^\*\*(.+?):\*\*\s*(.*)$", line)
        if m:
            flush()
            if m.group(1) not in labels:
                raise SpecError(f"{where}: unknown field '{m.group(1)}' (expected one of {', '.join(labels)})")
            key, buf = labels[m.group(1)], [m.group(2)]
        elif key:
            buf.append(line)
        elif line.strip():
            raise SpecError(f"{where}: text before the first **Field:** label: {line.strip()[:60]}")
    flush()
    return out


def _rows(blocks, start, fields, table_keys, title_key, parse_id, where):
    """Reads one table and the ### sections after it into dicts."""
    level, head, body = blocks[start]
    by_header = {f[2]: f[0] for f in fields if f[2]}
    labels = {f[3]: f[0] for f in fields if f[3]}
    header, rows = _table(body)
    if header is None:
        return []
    expected = [next(f[2] for f in fields if f[0] == k) for k in table_keys]
    if header != expected:
        raise SpecError(f"{where}: table columns {header} should be {expected}")
    items = []
    for r in rows:
        if len(r) != len(header):
            raise SpecError(f"{where}: row {r[:1]} has {len(r)} cells, expected {len(header)}")
        item = {f[0]: "" for f in fields}
        for h, v in zip(header, r):
            item[by_header[h]] = v
        item[fields[0][0]] = parse_id(item[fields[0][0]], where)
        items.append(item)
    sections = {}
    i = start + 1
    while i < len(blocks) and blocks[i][0] == 3:
        _, h3, b3 = blocks[i]
        key, _, title = h3.partition(" ")
        key = parse_id(key, where)
        if key in sections:
            raise SpecError(f"{where}: two sections for {h3}")
        sections[key] = (title.strip(), _fields(b3, labels, f"{where} / {h3}"))
        i += 1
    for item in items:
        k = item[fields[0][0]]
        if k not in sections:
            raise SpecError(f"{where}: no ### section for {display_id(k)}")
        title, values = sections.pop(k)
        if title != item[title_key]:
            raise SpecError(f"{where}: section heading '{display_id(k)} {title}' doesn't match the table ('{item[title_key]}')")
        item.update(values)
    if sections:
        raise SpecError(f"{where}: sections without a table row: {', '.join(display_id(k) for k in sections)}")
    return items


def display_id(k):
    return f"OQ-{k}" if isinstance(k, int) else k


def _req_id(v, where):
    if not ID_RE.match(v):
        raise SpecError(f"{where}: bad requirement id '{v}' (expected AB-123)")
    return v


def _q_id(v, where):
    m = re.match(r"^OQ-(\d+)$", v)
    if not m:
        raise SpecError(f"{where}: bad question id '{v}' (expected OQ-<n>)")
    return int(m.group(1))


def read_sheet_md(path):
    blocks = _blocks(path.read_text(encoding="utf-8"))
    where = path.relative_to(ROOT).as_posix()
    if len(blocks) < 2 or blocks[1][0] != 1:
        raise SpecError(f"{where}: must start with '# <sheet name>'")
    sheet = Sheet(blocks[1][1], path.name, _paragraphs(blocks[1][2]))
    if len(sheet.intro) != 2:
        raise SpecError(f"{where}: expected two paragraphs under the title (pillar line, where it stands)")
    for i, (level, head, body) in enumerate(blocks):
        if level == 2 and head == "Requirements":
            sheet.reqs = _rows(blocks, i, REQ_FIELDS, REQ_TABLE, "feature", _req_id, f"{where} / Requirements")
        elif level == 2 and head == "Decisions & open questions":
            sheet.questions = _rows(blocks, i, Q_FIELDS, Q_TABLE, "question", _q_id, f"{where} / Decisions")
        elif level == 2:
            raise SpecError(f"{where}: unexpected section '## {head}'")
    return sheet


def read_spec_md():
    readme = REQ_DIR / "README.md"
    blocks = _blocks(readme.read_text(encoding="utf-8"))
    spec = Spec(title=blocks[1][1], intro=_paragraphs(blocks[1][2]))
    for level, head, body in blocks[2:]:
        if level != 2:
            continue
        if head == "Mechanics":
            for line in body:
                m = re.match(r"^- \[(.+?)\]\((.+?)\)", line.strip())
                if m:
                    sheet = read_sheet_md(REQ_DIR / m.group(2))
                    if sheet.name != m.group(1):
                        raise SpecError(f"README.md: link '{m.group(1)}' points at a file titled '{sheet.name}'")
                    spec.sheets.append(sheet)
        elif head == "Design pillars":
            header, rows = _table(body)
            spec.pillars = [tuple(r) for r in rows]
        elif head != "Summary":
            spec.sections.append((head, [l.strip()[2:] for l in body if l.strip().startswith("- ")]))
    gloss = _blocks((REQ_DIR / "glossary.md").read_text(encoding="utf-8"))
    spec.glossary_intro = _paragraphs(gloss[1][2])
    spec.glossary = [(h, "\n\n".join(_paragraphs(b))) for lvl, h, b in gloss if lvl == 3]
    validate(spec)
    return spec


def validate(spec):
    errors, ids, qnums = [], {}, {}
    for s in spec.sheets:
        for r in s.reqs:
            if r["id"] in ids:
                errors.append(f"{r['id']} is in both {ids[r['id']]} and {s.file}")
            ids[r["id"]] = s.file
        for q in s.questions:
            if q["num"] in qnums:
                errors.append(f"OQ-{q['num']} is in both {qnums[q['num']]} and {s.file}")
            qnums[q["num"]] = s.file
    pillars = {p[0] for p in spec.pillars}
    for s in spec.sheets:
        for r in s.reqs:
            for key, _, _, _, allowed in REQ_FIELDS:
                if allowed and r[key] and r[key] not in allowed:
                    errors.append(f"{s.file} {r['id']}: {key} '{r[key]}' is not one of {', '.join(allowed)}")
            want = BUILT_FOR_STATUS.get(r["status"], "No")
            if r["built"] != want:
                errors.append(f"{s.file} {r['id']}: status {r['status']} needs Built '{want}', not '{r['built']}'")
            for dep in filter(None, (d.strip() for d in r["depends"].split(","))):
                if dep not in ids:
                    errors.append(f"{s.file} {r['id']}: depends on unknown {dep}")
            for p in filter(None, (p.strip() for p in r["pillar"].split("/"))):
                if p not in pillars:
                    errors.append(f"{s.file} {r['id']}: unknown pillar {p}")
        for q in s.questions:
            if q["state"] and q["state"] not in STATE:
                errors.append(f"{s.file} OQ-{q['num']}: state '{q['state']}' is not one of {', '.join(STATE)}")
            for rel in filter(None, (x.strip() for x in q["related"].split(","))):
                if rel not in ids:
                    errors.append(f"{s.file} OQ-{q['num']}: related to unknown {rel}")
    if errors:
        raise SpecError("\n".join(errors))


# ---------------------------------------------------------------- summary

def tally(spec):
    """The overview numbers, shared by the README summary and overview.xlsx."""
    def counts(reqs):
        n = sum(r["priority"] in PRIORITY for r in reqs)
        done = sum(r["status"] == "Done" for r in reqs)
        return ([sum(r["priority"] == p for r in reqs) for p in PRIORITY]
                + [n, done, sum(r["status"] == "In Progress" for r in reqs), done / n if n else 0,
                   sum(r["flag"] == "Conflict" for r in reqs), sum(r["flag"] == "Review" for r in reqs)])

    reqs = [r for s in spec.sheets for r in s.reqs]
    total = counts(reqs)
    return {
        "mechanics": [(s, counts(s.reqs)) for s in spec.sheets],
        "total": total,
        "status": [(st, c, c / total[4] if total[4] else 0)
                   for st in STATUS_ORDER for c in [sum(r["status"] == st for r in reqs)]],
        "flags": [(f, sum(r["flag"] == f for r in reqs)) for f in FLAG],
        "pillars": [(p, name, statement, len(rs), sum(r["priority"] == "Must" for r in rs),
                     sum(r["status"] == "Done" for r in rs))
                    for p, name, statement in spec.pillars for rs in [[r for r in reqs if r["pillar"] == p]]],
    }


MECH_HEADERS = ["Mechanic"] + PRIORITY + ["Total", "Done", "In Progress", "% Done", "Conflict", "Review"]


def _cell(v):
    return str(v).replace("|", "\\|")


def _md_table(headers, rows):
    out = ["| " + " | ".join(headers) + " |", "|" + "|".join("---" for _ in headers) + "|"]
    out += ["| " + " | ".join(_cell(c) for c in r) + " |" for r in rows]
    return out


def summary_md(spec):
    t = tally(spec)
    pct = lambda x: f"{round(100 * x)}%"
    out = ["<!-- BEGIN generated summary: tools/specs.py build rewrites this block -->", "",
           "### Requirements by mechanic", ""]
    rows = [[f"[{s.name}]({s.file})"] + c[:7] + [pct(c[7])] + c[8:] for s, c in t["mechanics"]]
    rows.append(["**Total**"] + t["total"][:7] + [pct(t["total"][7])] + t["total"][8:])
    out += _md_table(MECH_HEADERS, rows)
    out += ["", "### Requirements by status", ""]
    out += _md_table(["Status", "Count", "% of Total"], [[st, c, pct(p)] for st, c, p in t["status"]])
    out += ["", "### Audit flags", ""]
    out += _md_table(["Flag", "Count"], t["flags"])
    out += ["", "### Design pillar coverage", ""]
    out += _md_table(["Pillar", "Name", "Reqs", "Must", "Done"], [[p[0], p[1]] + list(p[3:]) for p in t["pillars"]])
    out += ["", "<!-- END generated summary -->"]
    return "\n".join(out)


def update_summary(spec):
    readme = REQ_DIR / "README.md"
    text = readme.read_text(encoding="utf-8")
    new = re.sub(r"<!-- BEGIN generated.*?<!-- END generated summary -->", lambda _: summary_md(spec), text, flags=re.S)
    if new != text:
        readme.write_text(new, encoding="utf-8", newline="\n")
        return True
    return False


# ---------------------------------------------------------------- model -> workbooks

DARK, LINK, GREY, TOTAL = "1F2A36", "1F4E79", "667085", "F2F4F7"
COLOURS = {"Must": "F4CCCC", "Should": "FCE5CD", "Could": "D9EAD3", "Won't": "E7E6E6", "Done": "B6D7A8",
           "Conflict": "F4CCCC", "Review": "FCE5CD", "Updated": "CFE2F3", "New": "D9EAD3",
           "Yes": "B6D7A8", "Partial": "FFF2CC", "No": "EFEFEF"}
REQ_WIDTHS = [9, 24, 52, 44, 52, 9, 11, 40, 10, 12, 11, 14, 7, 16, 12]
OV_WIDTHS = [26, 22, 9, 9, 9, 11, 9, 11, 9, 9, 9]
PERCENT = "0%;\\-0%;\\-"
WRAP = Alignment(wrap_text=True, vertical="top")


def _fill(colour):
    return PatternFill("solid", fgColor=colour)


def _header(ws, row, names, cols=None):
    for i, name in enumerate(names):
        c = ws.cell(row, cols[i] if cols else i + 1, name)
        c.font, c.fill, c.alignment = Font(size=10, bold=True, color="FFFFFF"), _fill(DARK), WRAP
    ws.row_dimensions[row].height = 22


def _put(ws, row, col, value, **font):
    c = ws.cell(row, col, value if value != "" else None)
    c.font, c.alignment = Font(size=10, **font), WRAP
    return c


def _title(ws, text):
    ws["A1"] = text
    ws["A1"].font = Font(size=16, bold=True, color=DARK)


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _lines(text, chars):
    return max(1, -(-len(text) // chars))


def _book(title):
    wb = openpyxl.Workbook()
    wb.active.title = title[:31]
    return wb, wb.active


def build_overview(spec):
    wb, ov = _book("Overview")
    _widths(ov, OV_WIDTHS)
    _title(ov, spec.title)
    row = 2
    for i, para in enumerate(spec.intro):
        ov.cell(row, 1, para).font = Font(size=10, color=GREY if i else None)
        row += 1
    for head, bullets in spec.sections:
        row += 1
        ov.cell(row, 1, head).font = Font(size=10, bold=True)
        for b in bullets:
            row += 1
            _put(ov, row, 1, "• " + b)
            ov.merge_cells(start_row=row, start_column=1, end_row=row, end_column=11)
            ov.row_dimensions[row].height = 18 if len(b) < 150 else 32
        row += 1

    t = tally(spec)
    row += 1
    ov.cell(row, 1, "Requirements by mechanic (each opens its own workbook)").font = Font(size=10, bold=True)
    row += 1
    _header(ov, row, MECH_HEADERS)
    for s, c in t["mechanics"] + [(None, t["total"])]:
        row += 1
        cell = _put(ov, row, 1, s.name if s else "Total", bold=not s, color=LINK if s else None)
        if s:
            cell.hyperlink = book_name(s.file)
        for col, v in enumerate(c, 2):
            _put(ov, row, col, v, bold=not s).number_format = PERCENT if col == 9 else "General"
        if not s:
            for col in range(1, 12):
                ov.cell(row, col).fill = _fill(TOTAL)

    row += 2
    ov.cell(row, 1, "Requirements by status").font = Font(size=10, bold=True)
    ov.cell(row, 6, "Audit flags").font = Font(size=10, bold=True)
    row += 1
    _header(ov, row, ["Status", "Count", "% of Total", "Flag", "Count"], [1, 2, 3, 6, 7])
    for i, (st, c, p) in enumerate(t["status"]):
        r = row + 1 + i
        _put(ov, r, 1, st)
        _put(ov, r, 2, c)
        _put(ov, r, 3, p).number_format = PERCENT
        if i < len(t["flags"]):
            _put(ov, r, 6, t["flags"][i][0]).fill = _fill(COLOURS[t["flags"][i][0]])
            _put(ov, r, 7, t["flags"][i][1])
    row += len(t["status"]) + 2

    ov.cell(row, 1, "Design pillar coverage (from the original pitch)").font = Font(size=10, bold=True)
    row += 1
    _header(ov, row, ["Pillar", "Name", "Pitch Statement", "Reqs", "Must", "Done"], [1, 2, 3, 9, 10, 11])
    ov.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
    for p in t["pillars"]:
        row += 1
        for col, v in zip([1, 2, 3, 9, 10, 11], p):
            _put(ov, row, col, v)
        ov.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
        ov.row_dimensions[row].height = 26

    row += 2
    ov.cell(row, 1, "See also").font = Font(size=10, bold=True)
    _put(ov, row + 1, 1, "Glossary", color=LINK).hyperlink = GLOSSARY_BOOK
    return wb


def build_mechanic(s):
    wb, ws = _book(s.name)
    _widths(ws, REQ_WIDTHS)
    _title(ws, s.name)
    ws["A2"] = s.intro[0]
    ws["A2"].font = Font(size=10)
    _put(ws, 3, 1, s.intro[1])
    ws.merge_cells("A3:G3")
    ws.row_dimensions[3].height = 14 * _lines(s.intro[1], 190) + 14
    audited = re.search(r"\((\d{4}-\d\d-\d\d)\)", s.intro[1])
    _header(ws, 5, [f"{f[1]} ({audited.group(1)})" if f[0] == "as_built" and audited else f[1] for f in REQ_FIELDS])
    for r, req in enumerate(s.reqs, 6):
        for c, f in enumerate(REQ_FIELDS, 1):
            _put(ws, r, c, req[f[0]])
    end = 5 + max(1, len(s.reqs))
    ws.freeze_panes = "C6"
    ws.auto_filter.ref = f"A5:{get_column_letter(len(REQ_FIELDS))}{end}"
    col = {f[0]: get_column_letter(i) for i, f in enumerate(REQ_FIELDS, 1)}
    for key, allowed in [("priority", PRIORITY), ("type", TYPE), ("status", STATUS), ("milestone", MILESTONE),
                         ("flag", FLAG), ("built", BUILT)]:
        dv = DataValidation(type="list", formula1='"' + ",".join(allowed) + '"', allow_blank=True)
        dv.add(f"{col[key]}6:{col[key]}{end}")
        ws.add_data_validation(dv)
    for key, values in [("priority", PRIORITY), ("status", ["Done"]), ("flag", FLAG), ("built", BUILT)]:
        L = col[key]
        for v in values:
            ws.conditional_formatting.add(f"{L}6:{L}{end}", FormulaRule(
                formula=[f'${L}6="{v}"'], fill=PatternFill(bgColor=COLOURS[v])))

    row = end + 3
    ws.cell(row, 1, "Decisions & open questions").font = Font(size=12, bold=True, color=DARK)
    row += 1
    _header(ws, row, [f[1] for f in Q_FIELDS])
    if not s.questions:
        _put(ws, row + 1, 1, "None yet.")
    for r, q in enumerate(s.questions, row + 1):
        for c, f in enumerate(Q_FIELDS, 1):
            _put(ws, r, c, q[f[0]])
    if s.questions:
        dv = DataValidation(type="list", formula1='"' + ",".join(STATE) + '"', allow_blank=True)
        dv.add(f"F{row + 1}:F{row + len(s.questions)}")
        ws.add_data_validation(dv)
    return wb


def build_glossary(spec):
    wb, ws = _book("Glossary")
    _widths(ws, [26, 110])
    _title(ws, "Glossary")
    for i, para in enumerate(spec.glossary_intro, 2):
        ws.cell(i, 1, para).font = Font(size=10)
    row = 3 + len(spec.glossary_intro)
    _header(ws, row, ["Term", "Definition"])
    for term, definition in spec.glossary:
        row += 1
        _put(ws, row, 1, term, bold=True)
        _put(ws, row, 2, definition)
        ws.row_dimensions[row].height = 14 * _lines(definition, 120) + 4
    ws.freeze_panes = f"A{3 + len(spec.glossary_intro) + 1}"
    return wb


def book_name(md_file):
    return Path(md_file).stem + ".xlsx"


def build_books(spec):
    """Every workbook, by file name: one per Markdown file."""
    books = {OVERVIEW_BOOK: build_overview(spec), GLOSSARY_BOOK: build_glossary(spec)}
    for s in spec.sheets:
        books[book_name(s.file)] = build_mechanic(s)
    return books


def deterministic_bytes(wb):
    """Saves with fixed timestamps so an unchanged spec gives an identical file."""
    wb.properties.created = wb.properties.modified = datetime.datetime(2026, 1, 1)
    raw = io.BytesIO()
    wb.save(raw)
    out = io.BytesIO()
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(created|modified)[^>]*>)[^<]*", rb"\g<1>2026-01-01T00:00:00Z", data)
            dst.writestr(zipfile.ZipInfo(info.filename, (1980, 1, 1, 0, 0, 0)), data, zipfile.ZIP_DEFLATED)
    return out.getvalue()


# ---------------------------------------------------------------- workbooks -> model (for diff)

def _s(v):
    return "" if v is None else str(v)


def _col_a(ws):
    return [ws.cell(r, 1).value for r in range(1, ws.max_row + 2)]


def _read_table(ws, head, fields):
    cols = {}
    for c in range(1, ws.max_column + 1):
        h = _s(ws.cell(head, c).value)
        for key, header, *_ in fields:
            if h == header or h.startswith(header + " ("):
                cols[key] = c
    items = []
    for r in range(head + 1, ws.max_row + 1):
        if ws.cell(r, 1).value is None:
            break
        item = {key: _s(ws.cell(r, c).value) for key, c in cols.items()}
        first = fields[0][0]
        v = ws.cell(r, cols[first]).value
        item[first] = v if isinstance(v, int) else _s(v)
        items.append(item)
    return items


def read_spec_xlsx(md_spec):
    """Reads the workbooks back into a spec, using the Markdown spec for the list of files."""
    spec = Spec()
    ov = openpyxl.load_workbook(XLSX_DIR / OVERVIEW_BOOK).active
    spec.title = _s(ov["A1"].value)
    for r in range(2, ov.max_row + 1):
        text, bold = ov.cell(r, 1).value, ov.cell(r, 1).font.b
        if text and text.startswith("Requirements by mechanic"):
            break
        if text and bold:
            spec.sections.append((text, []))
        elif text and spec.sections:
            spec.sections[-1][1].append(text.removeprefix("• "))
        elif text:
            spec.intro.append(text)
    start = _col_a(ov).index("Pillar") + 1
    for r in range(start + 1, ov.max_row + 1):
        if not ov.cell(r, 1).value:
            break
        spec.pillars.append(tuple(_s(ov.cell(r, c).value) for c in (1, 2, 3)))

    for m in md_spec.sheets:
        ws = openpyxl.load_workbook(XLSX_DIR / book_name(m.file)).active
        col_a = _col_a(ws)
        head = col_a.index("ID") + 1
        sheet = Sheet(_s(ws["A1"].value), m.file, [_s(ws.cell(r, 1).value) for r in range(2, head) if ws.cell(r, 1).value])
        sheet.reqs = _read_table(ws, head, REQ_FIELDS)
        sheet.questions = [q for q in _read_table(ws, col_a.index("#") + 1, Q_FIELDS) if isinstance(q["num"], int)]
        spec.sheets.append(sheet)

    gl = openpyxl.load_workbook(XLSX_DIR / GLOSSARY_BOOK).active
    head = _col_a(gl).index("Term") + 1
    spec.glossary_intro = [_s(gl.cell(r, 1).value) for r in range(2, head) if gl.cell(r, 1).value]
    spec.glossary = [(_s(gl.cell(r, 1).value), _s(gl.cell(r, 2).value))
                     for r in range(head + 1, gl.max_row + 1) if gl.cell(r, 1).value]
    return spec


def compare(a, b, a_name, b_name):
    """Lists every difference between two specs, field by field."""
    out = []

    def diff(where, x, y):
        if x != y:
            out.append(f"{where}\n    {a_name}: {x!r}\n    {b_name}: {y!r}")

    diff("Overview title", a.title, b.title)
    diff("Overview intro", a.intro, b.intro)
    diff("Overview sections", a.sections, b.sections)
    diff("Design pillars", a.pillars, b.pillars)
    diff("Glossary intro", a.glossary_intro, b.glossary_intro)
    diff("Glossary", a.glossary, b.glossary)
    for sa, sb in zip(a.sheets, b.sheets):
        diff(f"{sa.file} title", sa.name, sb.name)
        diff(f"{sa.name} intro", sa.intro, sb.intro)
        for kind, la, lb, key in [("requirement", sa.reqs, sb.reqs, "id"), ("question", sa.questions, sb.questions, "num")]:
            ma, mb = {x[key]: x for x in la}, {x[key]: x for x in lb}
            diff(f"{sa.name} {kind} list", list(ma), list(mb))
            for k in ma.keys() & mb.keys():
                for f in ma[k]:
                    diff(f"{sa.name} {display_id(k)} {f}", ma[k][f], mb[k].get(f, ""))
    return out


# ---------------------------------------------------------------- commands

def _hashes():
    return json.loads(HASHES.read_text()) if HASHES.exists() else {}


def _sha(data):
    return hashlib.sha256(data).hexdigest()


def _edited(hashes):
    """Workbooks we generated that were changed since."""
    return [n for n, h in hashes.items() if (XLSX_DIR / n).exists() and _sha((XLSX_DIR / n).read_bytes()) != h]


def cmd_build(force):
    spec = read_spec_md()
    hashes = _hashes()
    edited = _edited(hashes)
    if edited and not force:
        raise SpecError(f"edited since the last build: {', '.join(edited)}. Run 'specs.py diff', copy the changes into "
                        "the Markdown, then 'specs.py build --force'.")
    XLSX_DIR.mkdir(parents=True, exist_ok=True)
    written, new_hashes = [], {}
    for name, wb in build_books(spec).items():
        data, path = deterministic_bytes(wb), XLSX_DIR / name
        if not path.exists() or path.read_bytes() != data:
            path.write_bytes(data)
            written.append(name)
        new_hashes[name] = _sha(data)
    removed = [n for n in hashes if n not in new_hashes and (XLSX_DIR / n).exists()]
    for n in removed:
        (XLSX_DIR / n).unlink()
    HASHES.write_text(json.dumps(new_hashes, indent=2, sort_keys=True) + "\n")
    readme = update_summary(spec)
    print(f"{len(new_hashes)} workbooks, {len(written)} written"
          + (f" ({', '.join(written)})" if written else "")
          + (f"; removed {', '.join(removed)}" if removed else "")
          + ("; README summary updated" if readme else ""))


def cmd_diff():
    edited = _edited(_hashes())
    if not edited:
        print("no workbook was edited since the last build")
        return
    print(f"edited: {', '.join(edited)}")
    changes = compare(read_spec_md(), read_spec_xlsx(read_spec_md()), "markdown", "workbook")
    print(f"{len(changes)} difference(s)" + ("" if changes else " (formatting only)"))
    for c in changes:
        print(c)


def main(argv):
    try:
        if argv[:1] == ["build"]:
            cmd_build("--force" in argv)
        elif argv[:1] == ["diff"]:
            cmd_diff()
        else:
            print(__doc__)
            return 2
    except SpecError as e:
        print(f"error: {e}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
